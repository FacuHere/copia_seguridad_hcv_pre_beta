from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"docs\Presupuesto_estimativo_HCV_2026-10-22.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)

NAVY = "24485A"
PALE = "EAF1F4"
PALE2 = "F5F7F8"
GRID = "C9D3D8"
INK = "111111"
GREY = "4C5960"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run(run, size=9.4, bold=False, color=INK, italic=False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_text(cell, text, size=8.8, bold=False, color=INK, align=None):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    if align is not None:
        p.alignment = align
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break()
        r = p.add_run(line)
        set_run(r, size=size, bold=bold if i == 0 else False, color=color)
    set_cell_margins(cell)


def setup_table(table, widths, header=True, font_size=8.7):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        prevent_row_split(row)
        for idx, cell in enumerate(row.cells):
            if idx < len(widths):
                cell.width = Inches(widths[idx])
    if header and table.rows:
        repeat_table_header(table.rows[0])
        for cell in table.rows[0].cells:
            set_cell_shading(cell, NAVY)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    set_run(run, size=font_size, bold=True, color=WHITE)
    set_table_borders(table)


def add_table(doc, headers, rows, widths, font_size=8.7, right_cols=()):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, size=font_size, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_shading(table.rows[0].cells[i], NAVY)
    for r_idx, values in enumerate(rows):
        cells = table.add_row().cells
        if r_idx % 2:
            for cell in cells:
                set_cell_shading(cell, PALE2)
        for c_idx, value in enumerate(values):
            align = WD_ALIGN_PARAGRAPH.RIGHT if c_idx in right_cols else None
            set_cell_text(cells[c_idx], value, size=font_size, align=align)
    setup_table(table, widths, header=True, font_size=font_size)
    return table


def add_para(doc, text="", size=9.6, bold=False, color=INK, italic=False, before=0, after=5, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, color=color, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run(r, size=13 if level == 1 else 10.5, bold=True, color=INK)
    return p


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Inches(0.2 + level * 0.18)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(text)
    set_run(r, size=9.0)
    return p


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "176B87")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    r_pr.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    r_pr.append(size)
    new_run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def money(n):
    return "ARS " + f"{int(n):,}".replace(",", ".")


base_rows = [
    ("01  Relevamiento funcional y arquitectura\nEsfuerzo estimado para análisis del flujo web, arquitectura Express/MariaDB, reglas observables y límites del alcance.", 60, 45000, "Existe; horas no registradas"),
    ("02  Modelo de datos y migraciones\nDiseño relacional y scripts SQL 001-006 para cuentas, RBAC, pacientes, estudios, documentos, auditoría y journal.", 56, 60000, "Scripts presentes; DB no verificada"),
    ("03  Autenticación, sesiones y RBAC\nLogin/logout, sesiones persistentes, hashes, límite de intentos, middleware y herramientas CLI de administración inicial.", 88, 65000, "Código presente; runtime no verificado"),
    ("04  API clínica de pacientes y estudios\nLectura, creación y edición con validaciones, auditoría, control de versión y fecha de estudio inmutable.", 68, 65000, "Parcial; baja clínica ausente"),
    ("05  Interfaz clínica y conexión con API\nPantallas estáticas, formularios y flujos web para pacientes/estudios; tablas auxiliares siguen locales al navegador.", 56, 60000, "Código presente; alcance local parcial"),
    ("06  Ciclo de vida documental\nCarga streaming, validación, SHA-256, asociación paciente/estudio, listado, descarga, vista PDF/imagen, reemplazo y eliminación.", 130, 65000, "Implementación presente; fallo manual informado"),
    ("07  Journal, recuperación y auditoría documental\nEstados de operación, verificación de archivos, recuperación de cargas/reemplazos/bajas/movimientos y manejo de staging.", 76, 65000, "Código presente; recuperación real no verificada"),
    ("08  Pruebas automatizadas y fixtures\nPruebas de contratos, seguridad, frontend, rutas HTTP, filesystem temporal y adaptadores MariaDB simulados.", 84, 50000, "52 casos declarados; no ejecutados aquí"),
    ("09  Entorno de validación Docker\nCompose de prueba con MariaDB, API, phpMyAdmin y volumen de archivos separado; no equivale a despliegue productivo.", 32, 48000, "Configuración presente; operación no validada ahora"),
    ("10  Documentación funcional y técnica\nREADME, guía SQL/migraciones y documento con 27 casos de uso, incluidos los límites pendientes.", 48, 45000, "Documentos presentes; sin acta de aceptación"),
]

future_rows = [
    ("11  Cerrar decisiones funcionales pendientes\nDefinir campos por servicio, alcance de permisos, retención/borrado, tamaños de archivo y criterios institucionales de aceptación.", 24, 45000, "Pendiente"),
    ("12  Diagnosticar y corregir la carga documental informada\nInvestigar el caso sin registro en MariaDB/filesystem, revisar la discrepancia de formatos en la interfaz y corregir solo con evidencia.", 32, 65000, "Pendiente; causa no determinada"),
    ("13  Verificación dirigida de carga sobre entorno aislado\nEjecutar caso sintético con MariaDB real y filesystem; contrastar fila, asociación, ruta, hash, descarga y rechazo de formatos/tamaño.", 24, 50000, "Pendiente"),
    ("14  Administración web de usuarios y roles - backend\nRutas protegidas, validaciones, asignaciones RBAC, auditoría y pruebas de seguridad para administración autorizada.", 40, 65000, "No existe CRUD web suficiente"),
    ("15  Administración web de usuarios y roles - frontend\nPantallas de altas, modificación/desactivación, asignación de roles y mensajes de operación.", 32, 60000, "No existe UI suficiente"),
    ("16  Backup y restauración coordinados\nDiseño e implementación de copia conjunta MariaDB + raíz física, manifiesto/verificación de hashes y procedimiento de restauración ensayable.", 96, 65000, "Pendiente; requiere definición de retención/medio"),
    ("17  Pruebas integrales de regresión y recuperación\nValidar autenticación, RBAC, pacientes, estudios, documentos, reinicio/fallos controlados y restauración en entorno descartable.", 48, 50000, "Pendiente; suite no ejecutada en este análisis"),
    ("18  Preparación de instalación y transferencia técnica\nGuía operativa de despliegue, TLS/red, permisos de almacenamiento, migraciones, cuentas, backup/restore y entrega técnica.", 32, 48000, "Pendiente; no incluye instalación productiva"),
]

base_total = sum(hours * rate for _, hours, rate, _ in base_rows)
future_total = sum(hours * rate for _, hours, rate, _ in future_rows)
assert base_total == 40_846_000
assert future_total == 19_056_000
assert base_total + future_total == 59_902_000

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
sec.top_margin = Inches(0.62)
sec.bottom_margin = Inches(0.62)
sec.left_margin = Inches(0.66)
sec.right_margin = Inches(0.66)
sec.header_distance = Inches(0.28)
sec.footer_distance = Inches(0.28)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Arial"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
normal.font.size = Pt(9.5)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(4)
for name in ("Heading 1", "Heading 2", "Title", "Subtitle"):
    style = styles[name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    style.font.color.rgb = RGBColor.from_string(INK)

header = sec.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.paragraph_format.space_after = Pt(0)
r = header.add_run("HISTORIAL CLÍNICO VISUAL  |  PRESUPUESTO ESTIMATIVO")
set_run(r, size=7.8, bold=True, color=GREY)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.paragraph_format.space_before = Pt(0)
r = footer.add_run("HCV  |  Referencia económica al 22/10/2026  |  Página ")
set_run(r, size=7.8, color=GREY)
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
footer._p.append(fld)

# Page 1 - formal cover and executive summary.
p = doc.add_paragraph(style="Title")
p.paragraph_format.space_before = Pt(14)
p.paragraph_format.space_after = Pt(4)
r = p.add_run("Presupuesto estimativo del sistema HCV")
set_run(r, size=22, bold=True, color=INK)
add_para(doc, "Desarrollo, integración y tareas de finalización de Historial Clínico Visual", size=12, color=GREY, after=12)

metadata = [
    ("Presupuesto Nº", "Pendiente de asignación"),
    ("Fecha de emisión", "22 de octubre de 2026"),
    ("Destinatario", "[Organismo público / unidad solicitante]"),
    ("Solicitado por / cargo", "[Completar]"),
    ("Proveedor / responsable", "[Completar]"),
    ("CUIT / condición fiscal", "[Completar y validar]"),
    ("Unidad de negocio", "Software - desarrollo e implementación"),
    ("Moneda", "Pesos argentinos (ARS), valores nominales al 22/10/2026"),
]
table = doc.add_table(rows=0, cols=2)
for label, value in metadata:
    cells = table.add_row().cells
    set_cell_text(cells[0], label, size=9.0, bold=True)
    set_cell_shading(cells[0], PALE)
    set_cell_text(cells[1], value, size=9.0)
setup_table(table, [1.75, 5.16], header=False)

add_heading(doc, "Resumen económico", 1)
summary_rows = [
    ("Valoración referencial de componentes y documentación existentes", money(base_total)),
    ("Estimación de tareas pendientes para el alcance de cierre descrito", money(future_total)),
    ("TOTAL INDICATIVO DEL ALCANCE VALORADO", money(base_total + future_total)),
]
sum_table = doc.add_table(rows=0, cols=2)
for i, (label, value) in enumerate(summary_rows):
    cells = sum_table.add_row().cells
    set_cell_text(cells[0], label, size=10 if i == 2 else 9.4, bold=(i == 2))
    set_cell_text(cells[1], value, size=11 if i == 2 else 10, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
    set_cell_shading(cells[0], NAVY if i == 2 else PALE)
    set_cell_shading(cells[1], NAVY if i == 2 else PALE2)
    if i == 2:
        for cell in cells:
            for p2 in cell.paragraphs:
                for run in p2.runs:
                    run.font.color.rgb = RGBColor.from_string(WHITE)
setup_table(sum_table, [4.85, 2.06], header=False)

add_para(doc, "Importes antes de impuestos: el tratamiento fiscal no se calcula porque dependen de la condición del proveedor y del procedimiento de contratación del organismo.", size=9.0, bold=True, color=GREY, before=6, after=5)
add_para(doc, "La cifra consolidada valora el código y los documentos observables, más una etapa separada de cierre. Las horas son estimaciones reconstruidas por paquetes de trabajo; no son partes horarios ni registros de dedicación. El alcance no acredita aceptación institucional ni equivale a una factura, adjudicación o liquidación de trabajo efectivamente realizado.", size=9.1, after=4)
add_para(doc, "Conclusión ejecutiva: existe una aplicación web cliente-servidor con módulos clínicos, autenticación/RBAC y código de almacenamiento documental recuperable. No se considera aún verificado el funcionamiento integral con MariaDB y archivos físicos: se informó una prueba manual en Docker en la que los documentos no quedaron listados ni registrados en la base. El presupuesto de cierre contempla investigar y resolver ese punto antes de darlo por aceptado.", size=9.1, after=4)

# Page 2 - scope and current evidence status.
doc.add_page_break()
add_heading(doc, "1. Alcance técnico y estado observado", 1)
add_para(doc, "La valoración se limita al repositorio HCV y a la documentación disponible al 22/10/2026. El relevamiento de código es estático; no se consultó una base de datos, no se inició la aplicación y no se ejecutó la suite durante esta preparación.", size=9.1, after=7)
scope_rows = [
    ("Autenticación y sesiones", "Login/logout, /api/auth/me, cookie de sesión y persistencia en MariaDB; hashes de contraseña y límite técnico de intentos.", "Código presente; ejecución no certificada"),
    ("Roles y permisos", "Middleware backend basado en user_roles -> role_permissions -> permissions; roles funcionales en migraciones.", "Modelo presente; gestión web incompleta"),
    ("Pacientes", "Búsqueda/listado, consulta, alta y edición con validación/auditoría.", "Baja no disponible; política pendiente"),
    ("Estudios", "Consulta, alta y edición asociados a paciente; study_date inmutable y reasignación de paciente rechazada.", "Borrado no disponible; reglas por servicio pendientes"),
    ("Documentos", "Código para cargar/listar/descargar/reemplazar/eliminar y verificar SHA-256.", "No aceptar como cerrado: prueba manual informó ausencia de filas/archivos"),
    ("Word / DOCX", "El backend reconoce .doc y .docx y valida estructura DOCX; la entrada inicial de carga del frontend no ofrece esos formatos en su atributo accept.", "Discrepancia de interfaz; investigación y prueba pendientes. Word no tiene vista en línea, solo descarga"),
    ("Journal y recuperación", "Journal de operaciones y recuperación al iniciar; la presencia del .service.lock residual se ignora.", "Implementado en código; recuperación real no verificada"),
    ("Backups y administración", "Exportación/importación JSON cubre datos locales del navegador; no incluye la base ni los documentos físicos. Bootstrap/reset de cuenta son CLI.", "Backup coordinado y CRUD web de usuarios/roles pendientes"),
    ("Despliegue y seguridad operativa", "Compose de pruebas con MariaDB, API, phpMyAdmin y volumen privado de documentos.", "Entorno de prueba; TLS, despliegue productivo y operación no verificados"),
]
add_table(doc, ["Módulo", "Evidencia y alcance del código", "Estado para presupuesto"], scope_rows, [1.28, 3.95, 1.68], font_size=8.2)
add_para(doc, "El inventario funcional disponible contiene 27 casos de uso: 19 descritos como completos a nivel de flujo de código, 2 completos con alcance limitado y 6 parciales. En ese documento, “completo” no significa probado contra MariaDB real ni aceptado por un organismo.", size=8.8, color=GREY, before=6, after=4)
add_para(doc, "Fuera de alcance de esta estimación: PACS/DICOMweb, visor DICOM avanzado, validación semántica DICOM, firma digital, importación de datos clínicos reales y rediseño de procesos por servicio sin requisitos aprobados.", size=8.8, after=0)

# Page 3 - existing implementation valuation.
doc.add_page_break()
add_heading(doc, "2. Valoración del desarrollo observable", 1)
add_para(doc, "Las cantidades expresan horas profesionales estimadas para producir los componentes que hoy se encuentran en código o documentación. El estado indica qué se puede comprobar en los archivos, no el resultado de una aceptación funcional.", size=9.0, after=6)
base_table_rows = []
for label, hours, rate, status in base_rows:
    base_table_rows.append((label, str(hours), money(rate), money(hours * rate), status))
base_table = add_table(doc, ["Actividad / entregable y evidencia resumida", "h", "Tarifa por h", "Subtotal", "Estado"], base_table_rows, [3.18, 0.40, 0.91, 1.16, 1.26], font_size=7.65, right_cols=(1, 2, 3))
cells = base_table.add_row().cells
merged = cells[0].merge(cells[2])
set_cell_text(merged, "Subtotal estimado del desarrollo observable", size=9.1, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
set_cell_text(cells[3], money(base_total), size=9.1, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
set_cell_text(cells[4], "698 h", size=8.7, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
for cell in cells:
    set_cell_shading(cell, PALE)
prevent_row_split(base_table.rows[-1])

add_heading(doc, "Alcance de la valoración", 2)
add_bullet(doc, "Incluye el código y la documentación existentes de la aplicación: interfaz, API, persistencia, seguridad, archivos, pruebas fuente, configuración de prueba y documentación de casos de uso.")
add_bullet(doc, "La implementación documental se valora como código presente, pero queda sujeta a diagnóstico: el usuario informó que una carga de prueba no apareció en la interfaz ni generó registro visible en MariaDB.")
add_bullet(doc, "Los tests fuente sirven como evidencia de diseño de pruebas; contienen adaptadores simulados, filesystem temporal y servidores HTTP locales. Esta estimación no afirma que hayan pasado ni que cubran MariaDB real.")
add_para(doc, "No se asigna valor como “funcionalidad terminada y aceptada” a los componentes no verificados; la tabla es una valoración económica del alcance visible y no una certificación de calidad.", size=8.9, bold=True, color=GREY, before=4, after=0)

# Page 4 - estimated closure work.
doc.add_page_break()
add_heading(doc, "3. Presupuesto de tareas de finalización", 1)
add_para(doc, "Las tareas siguientes son futuras y se muestran separadas de la valoración existente. El esfuerzo parte de la información disponible y debe reestimarse si el organismo agrega servicios, reglas o requisitos de infraestructura.", size=9.0, after=6)
future_table_rows = []
for label, hours, rate, status in future_rows:
    future_table_rows.append((label, str(hours), money(rate), money(hours * rate), status))
future_table = add_table(doc, ["Actividad / entregable esperado", "h", "Tarifa por h", "Subtotal", "Estado"], future_table_rows, [3.18, 0.40, 0.91, 1.16, 1.26], font_size=7.65, right_cols=(1, 2, 3))
cells = future_table.add_row().cells
merged = cells[0].merge(cells[2])
set_cell_text(merged, "Subtotal de tareas de finalización", size=9.1, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
set_cell_text(cells[3], money(future_total), size=9.1, bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)
set_cell_text(cells[4], "328 h", size=8.7, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
for cell in cells:
    set_cell_shading(cell, PALE)
prevent_row_split(future_table.rows[-1])

add_heading(doc, "Entregables verificables de cierre", 2)
add_bullet(doc, "Resolución documentada de la discrepancia de carga y evidencia concordante de API, MariaDB y filesystem sintético: asociación correcta, storage_key relativo, archivo privado y SHA-256 coincidente.")
add_bullet(doc, "Administración web de usuarios y roles limitada a permisos aprobados, con verificación de accesos permitidos y denegados.")
add_bullet(doc, "Copia y restauración coordinadas de base y archivos, comprobadas en un entorno descartable con inventario, hashes y relaciones verificadas.")
add_bullet(doc, "Acta o matriz de aceptación que separe pruebas aprobadas, fallidas y no ejecutadas; guía operativa para instalar y recuperar el sistema.")
add_para(doc, "Las horas de backup/restauración son provisionales hasta acordar frecuencia, retención, medio, cifrado, responsabilidades y RPO/RTO. No incluyen compra de almacenamiento ni operación recurrente.", size=8.8, color=GREY, before=4, after=0)

# Page 5 - totals, terms and exclusions.
doc.add_page_break()
add_heading(doc, "4. Resumen, inclusiones y condiciones", 1)
total_rows = [
    ("Valoración referencial del desarrollo observable", "698 h", money(base_total)),
    ("Tareas pendientes de cierre incluidas en esta estimación", "328 h", money(future_total)),
    ("TOTAL INDICATIVO DEL ALCANCE VALORADO", "1.026 h", money(base_total + future_total)),
]
total_table = doc.add_table(rows=1, cols=3)
for i, h in enumerate(("Componente", "Esfuerzo", "Importe")):
    set_cell_text(total_table.rows[0].cells[i], h, size=9, bold=True, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_shading(total_table.rows[0].cells[i], NAVY)
for idx, row in enumerate(total_rows):
    cells = total_table.add_row().cells
    for j, value in enumerate(row):
        set_cell_text(cells[j], value, size=9.2 if idx == 2 else 8.9, bold=idx == 2, align=WD_ALIGN_PARAGRAPH.RIGHT if j else None)
        if idx == 2:
            set_cell_shading(cells[j], PALE)
        elif idx % 2:
            set_cell_shading(cells[j], PALE2)
setup_table(total_table, [4.35, 0.88, 1.68], header=True, font_size=9)

add_heading(doc, "Incluido en el alcance valorado", 2)
for text in [
    "Desarrollo y documentación del software observado en el repositorio; tareas futuras acotadas a los rubros 11-18.",
    "Integración de aplicación Node.js/Express con MariaDB y almacenamiento privado en filesystem, en los límites evidenciados.",
    "Validaciones sintéticas y documentación técnica de operación previstas en la etapa de cierre.",
]:
    add_bullet(doc, text)

add_heading(doc, "No incluido o pendiente de definición", 2)
for text in [
    "Impuestos, tasas, retenciones o tratamiento fiscal; deben determinarse con CUIT, condición del proveedor y reglas del organismo.",
    "Servidor, alojamiento, dominio, certificados, licencias de terceros, red, hardware, energía, almacenamiento externo y costos de soporte continuo.",
    "Datos clínicos reales, migración/carga histórica, validación médico-funcional, procesos específicos aún sin especificación y borrado/retención de pacientes o estudios.",
    "PACS, DICOMweb, visor DICOM avanzado, firma digital, integraciones con equipamiento o sistemas externos.",
    "Servicio de guardia, mantenimiento evolutivo, actualizaciones posteriores y capacitación de usuarios finales; cotización separada según participantes y modalidad.",
]:
    add_bullet(doc, text)

add_heading(doc, "Condiciones comerciales a completar", 2)
terms = [
    ("Forma de pago", "A convenir conforme al procedimiento de contratación del organismo."),
    ("Validez propuesta", "15 días corridos desde la emisión, hasta el 06/11/2026; confirmar antes de presentar o contratar."),
    ("Plazo de entrega", "A definir después de aprobar requisitos y entorno de instalación."),
    ("Garantía", "Duración, alcance y exclusiones pendientes de acuerdo."),
    ("Impuestos", "No calculados; tratamiento fiscal pendiente de validación."),
    ("Especificación técnica", "Incluida en este documento como alcance, estimación y metodología."),
]
terms_table = doc.add_table(rows=0, cols=2)
for label, value in terms:
    cells = terms_table.add_row().cells
    set_cell_text(cells[0], label, size=8.7, bold=True)
    set_cell_shading(cells[0], PALE)
    set_cell_text(cells[1], value, size=8.7)
setup_table(terms_table, [1.45, 5.46], header=False)

# Page 6 - methodology, references, evidence and acceptance fields.
doc.add_page_break()
add_heading(doc, "5. Metodología y referencias económicas", 1)
add_para(doc, "Método: estimación ascendente por paquetes de trabajo. Para cada componente se estimó esfuerzo en horas-persona y se multiplicó por una tarifa horaria orientativa de perfil. No se reconstruyeron partes horarios ni se infirió dedicación a partir de la cantidad de líneas de código. Las horas describen una valuación económica del entregable visible y del alcance de cierre, no el tiempo real trabajado.", size=9.0, after=6)

rate_rows = [
    ("Analista funcional", money(45000), "CPCI San Juan, analista funcional senior: desde ARS 45.000/hora."),
    ("Desarrollo backend / integración", money(65000), "CPCI San Juan, backend: desde ARS 65.000/hora; contraste CIEC 2026."),
    ("Desarrollo frontend", money(60000), "CPCI San Juan, frontend: desde ARS 60.000/hora."),
    ("DBA / diseño de datos", money(60000), "CPCI San Juan, DBA: desde ARS 60.000/hora."),
    ("Testing / QA", money(50000), "CPCI San Juan, testing de aplicaciones: desde ARS 50.000/hora."),
    ("Implementación", money(48000), "CPCI San Juan, implementador de sistemas: desde ARS 48.000/hora."),
]
add_table(doc, ["Perfil usado", "Tarifa aplicada", "Referencia"], rate_rows, [1.66, 1.18, 4.07], font_size=8.1, right_cols=(1,))

add_heading(doc, "Lectura de las tarifas y límites", 2)
add_bullet(doc, "Se toman tarifas horarias profesionales publicadas como referencias orientativas; el CPCI no informa fecha explícita de actualización en la página consultada y representa una referencia provincial, no una tarifa obligatoria nacional.")
add_bullet(doc, "Como contraste, el CIEC publica para 2026 1 DTG = ARS 130.000; los coeficientes de analista programador y programador Linux equivalen aproximadamente a ARS 36.400-61.100/h y ARS 42.900-66.300/h. La tarifa backend aplicada se ubica dentro del segundo rango.")
add_bullet(doc, "Los salarios brutos mensuales no se convierten en tarifas de factura: por ejemplo, CESSI/OPSSI reportó ARS 2.725.272 como salario bruto promedio del sector a marzo de 2025. Ese dato no es una tarifa independiente ni se usó para calcular estos importes.")
add_bullet(doc, "No se agregaron cargas laborales, utilidad empresarial, riesgo financiero, inflación futura ni impuestos porque se desconoce si contratará un empleado, profesional independiente o empresa proveedora, y faltan los datos fiscales y contractuales.")

add_heading(doc, "Fuentes económicas consultadas", 2)
sources = [
    ("CPCI San Juan. Honorarios sugeridos. Valores horarios por perfil disponibles al 22/10/2026; fecha de actualización no indicada en la página.", "https://cpci.com.ar/Honorarios.html", "Abrir referencia CPCI"),
    ("Colegio de Ingenieros Especialistas de Córdoba (CIEC). Honorarios Informática 2026. Vigencia 01/01/2026-31/12/2026; 1 DTG = ARS 130.000.", "https://ciec.com.ar/archivos/HONORARIOS%20DTO%20INFORMATICA%202026.pdf", "Abrir tabla CIEC 2026"),
    ("CESSI/OPSSI. Reporte de indicadores de la industria del software con foco en talento, julio de 2025. Dato salarial bruto sectorial, usado solo para distinguir salarios de honorarios.", "https://cessi.org.ar/wp-content/uploads/2025/08/CESSI-OPSSI-Reporte-de-los-principales-indicadores-de-la-industria-del-Software-con-Foco-en-Talento_Julio-2025.pdf", "Abrir reporte CESSI/OPSSI"),
]
for desc, url, label in sources:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.02
    r = p.add_run(desc + " ")
    set_run(r, size=8.2)
    add_hyperlink(p, label, url)

add_heading(doc, "Evidencia del alcance revisado", 2)
add_para(doc, "Fuentes del proyecto: README.md; sql/README.md; sql/001_initial_schema.sql a sql/006_documents_storage_lifecycle.sql; src/server.js; src/routes/auth.js, patients.js, studies.js y documents.js; src/services/document-storage.js, document-operations.js y storage-readiness.js; public/index.html y public/js/app.js; 13 archivos *.test.js con 52 declaraciones test() visibles; docs/Casos de uso HCV.docx (27 casos).", size=8.3, after=4)
add_para(doc, "El historial Git accesible muestra un único commit base (cd71b69, 15/10/2026); no permite reconstruir tareas ni horas reales. La documentación de casos de uso indica relevamiento estático y que no confirma el estado de migraciones aplicadas en una instancia particular.", size=8.3, color=GREY, after=5)
add_para(doc, "Referencia de formato: el PDF Anexo E de ejemplo aporta los campos de identificación, detalle de servicios, alcance/exclusiones, forma de pago, validez, garantía, plazo, firmas y aceptación. Se omitieron del ejemplo Oracle, hardware y paquetes de mantenimiento por no corresponder al alcance de HCV.", size=8.3, after=5)

add_heading(doc, "Datos de emisión y aceptación", 2)
accept = doc.add_table(rows=3, cols=2)
for i, (left, right) in enumerate([
    ("Elaborado por: [Completar]", "Firma: [Completar]"),
    ("Aceptación del organismo", "Firma y aclaración: ______________________________"),
    ("Cargo / fecha", "______________________________________________"),
]):
    set_cell_text(accept.rows[i].cells[0], left, size=8.8, bold=(i == 0))
    set_cell_text(accept.rows[i].cells[1], right, size=8.8, bold=(i == 0))
    if i == 0:
        set_cell_shading(accept.rows[i].cells[0], PALE)
        set_cell_shading(accept.rows[i].cells[1], PALE)
setup_table(accept, [3.45, 3.46], header=False)

doc.core_properties.title = "Presupuesto estimativo del sistema HCV"
doc.core_properties.subject = "Valoración técnica y económica del desarrollo de Historial Clínico Visual"
doc.core_properties.author = ""
doc.core_properties.keywords = "HCV, presupuesto, ARS, desarrollo de software"
doc.save(OUT)
print(f"CREATED {OUT.resolve()}")
print(f"BASE {base_total} ARS / {sum(r[1] for r in base_rows)} h")
print(f"FUTURE {future_total} ARS / {sum(r[1] for r in future_rows)} h")
print(f"TOTAL {base_total + future_total} ARS / {sum(r[1] for r in base_rows) + sum(r[1] for r in future_rows)} h")
